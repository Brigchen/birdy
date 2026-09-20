# -*- coding: utf-8 -*-
"""CSV / Excel 导出与 matplotlib 柱状图。"""

from __future__ import annotations

import csv
import os
from pathlib import Path
from typing import Any, Dict, List, Sequence

_FONT_PROP = None
_FONT_READY = False


def wrap_group_label(text: str, max_chars: int = 18) -> str:
    """长分组名按「·」和字数折行，供需要多行的场合使用。"""
    text = str(text or "").strip()
    if not text:
        return text
    chunks: List[str] = []
    parts = text.split(" · ") if " · " in text else [text]
    for part in parts:
        part = part.strip()
        if not part:
            continue
        if len(part) <= max_chars:
            chunks.append(part)
            continue
        for i in range(0, len(part), max_chars):
            chunks.append(part[i : i + max_chars])
    return "\n".join(chunks) if chunks else text


def short_group_label(text: str, max_chars: int = 20) -> str:
    """单行截断标签（表格提示等）；柱状图改用 wrap_group_label 以显示全名。"""
    text = " ".join(str(text or "").split())
    if len(text) <= max_chars:
        return text
    return text[: max(1, max_chars - 1)] + "…"


def safe_export_stem(folder: str) -> str:
    name = Path(folder).name.strip() or "mtf"
    out = "".join("_" if c in '<>:"/\\|?*' else c for c in name).strip(" .")
    return (out or "mtf")[:60]


def suggested_xlsx_path(photo_folder: str, *, now=None) -> Path:
    """默认保存到照片目录：文件夹名_年月日_编号.xlsx。"""
    from datetime import datetime

    folder = Path(photo_folder) if photo_folder else Path(".")
    if not folder.is_dir():
        from .paths import default_output_dir

        folder = default_output_dir()
        folder.mkdir(parents=True, exist_ok=True)
    stem = safe_export_stem(str(folder))
    day = (now or datetime.now()).strftime("%Y%m%d")
    n = 1
    while n < 1000:
        p = folder / f"{stem}_{day}_{n:02d}.xlsx"
        if not p.exists():
            return p
        n += 1
    return folder / f"{stem}_{day}_{n:02d}.xlsx"


def _label_units(text: str) -> int:
    return sum(2 if ord(c) > 127 else 1 for c in text)


def cjk_font_prop():
    """matplotlib 中文字体（Windows 雅黑 / 黑体，macOS 苹方，Linux Noto）。"""
    global _FONT_PROP, _FONT_READY
    if _FONT_READY:
        return _FONT_PROP
    _FONT_READY = True
    try:
        from matplotlib import font_manager as fm
        from matplotlib import rcParams
    except Exception:
        return None
    win = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    files = [
        win / "msyh.ttc",
        win / "msyh.ttf",
        win / "msyhbd.ttc",
        win / "simhei.ttf",
        win / "simsun.ttc",
        Path("/System/Library/Fonts/PingFang.ttc"),
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"),
    ]
    chosen_name = None
    for fp in files:
        if not fp.is_file():
            continue
        try:
            fm.fontManager.addfont(str(fp))
            _FONT_PROP = fm.FontProperties(fname=str(fp))
            chosen_name = _FONT_PROP.get_name()
            break
        except Exception:
            continue
    if _FONT_PROP is None:
        for f in fm.fontManager.ttflist:
            n = f.name or ""
            if any(
                k in n
                for k in ("YaHei", "SimHei", "PingFang", "Noto Sans CJK", "Source Han")
            ):
                chosen_name = n
                _FONT_PROP = fm.FontProperties(family=n)
                break
    rcParams["axes.unicode_minus"] = False
    if chosen_name:
        rcParams["font.sans-serif"] = [chosen_name, "DejaVu Sans", "Arial"]
        rcParams["font.family"] = "sans-serif"
    return _FONT_PROP


def _set_text_font(artist, prop) -> None:
    if artist is None or prop is None:
        return
    try:
        artist.set_fontproperties(prop)
    except Exception:
        pass


def export_iso_csv(path: str, rows: List[Dict[str, Any]]) -> None:
    fields = [
        "file",
        "lens",
        "focal_mm",
        "iso",
        "aperture",
        "camera",
        "group",
        "n_edges",
        "mtf50_cy_px",
        "mtf50_center",
        "mtf50_corner",
        "mtf50_lp_mm",
        "mtf30_cy_px",
        "mtf10_cy_px",
        "mtf_peak",
        "center_edge_ratio",
        "tenengrad",
        "datetime",
        "path",
        "error",
    ]
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            meta = r.get("meta") or {}
            w.writerow(
                {
                    "file": meta.get("file"),
                    "lens": meta.get("lens"),
                    "focal_mm": meta.get("focal_mm"),
                    "iso": meta.get("iso"),
                    "aperture": meta.get("aperture"),
                    "camera": meta.get("camera"),
                    "group": r.get("group"),
                    "n_edges": r.get("n_edges"),
                    "mtf50_cy_px": _fmt(r.get("mtf50_cy_px")),
                    "mtf50_center": _fmt(r.get("mtf50_center")),
                    "mtf50_corner": _fmt(r.get("mtf50_corner")),
                    "mtf50_lp_mm": _fmt(r.get("mtf50_lp_mm")),
                    "mtf30_cy_px": _fmt(r.get("mtf30_cy_px")),
                    "mtf10_cy_px": _fmt(r.get("mtf10_cy_px")),
                    "mtf_peak": _fmt(r.get("mtf_peak")),
                    "center_edge_ratio": _fmt(r.get("center_edge_ratio"), 3),
                    "tenengrad": _fmt(r.get("tenengrad")),
                    "datetime": meta.get("datetime"),
                    "path": meta.get("path") or r.get("path"),
                    "error": r.get("error") or "",
                }
            )


def export_scene_csv(path: str, scenes: List[Dict[str, Any]]) -> None:
    fields = [
        "scene",
        "rank",
        "relative",
        "lens",
        "focal_mm",
        "iso",
        "aperture",
        "mtf50_cy_px",
        "mtf30_cy_px",
        "mtf10_cy_px",
        "mtf_peak",
        "center_edge_ratio",
        "tenengrad",
        "file",
        "path",
    ]
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        for sc in scenes:
            name = sc.get("name")
            for r in sc.get("ranked") or []:
                w.writerow(
                    {
                        "scene": name,
                        "rank": r.get("rank"),
                        "relative": _fmt(r.get("relative"), 3),
                        "lens": r.get("lens"),
                        "focal_mm": r.get("focal_mm"),
                        "iso": r.get("iso"),
                        "aperture": r.get("aperture"),
                        "mtf50_cy_px": _fmt(r.get("mtf50_cy_px")),
                        "mtf30_cy_px": _fmt(r.get("mtf30_cy_px")),
                        "mtf10_cy_px": _fmt(r.get("mtf10_cy_px")),
                        "mtf_peak": _fmt(r.get("mtf_peak")),
                        "center_edge_ratio": _fmt(r.get("center_edge_ratio"), 3),
                        "tenengrad": _fmt(r.get("tenengrad")),
                        "file": r.get("file"),
                        "path": r.get("path"),
                    }
                )


def _fmt(v: Any, nd: int = 4) -> str:
    if v is None:
        return ""
    try:
        fv = float(v)
    except (TypeError, ValueError):
        return str(v)
    if fv != fv:  # nan
        return ""
    return f"{fv:.{nd}f}"


def _num(v: Any, nd: int = 4):
    if v is None:
        return None
    try:
        fv = float(v)
    except (TypeError, ValueError):
        return None
    if fv != fv:
        return None
    return round(fv, nd)


def _xlsx_header(ws, titles: Sequence[str]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    fill = PatternFill("solid", fgColor="E8F5E9")
    font = Font(bold=True, color="1B5E20")
    for i, t in enumerate(titles, start=1):
        cell = ws.cell(1, i, t)
        cell.font = font
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _xlsx_widths(ws, widths: Sequence[float]) -> None:
    from openpyxl.utils import get_column_letter

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last = max(1, int(ws.max_row or 1))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{last}"


def export_iso_xlsx(path: str, rows: List[Dict[str, Any]]) -> None:
    from openpyxl import Workbook

    from .grouping import summarize_groups

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "文件明细"
    file_headers = [
        "文件",
        "镜头",
        "焦距",
        "ISO",
        "光圈",
        "机身",
        "分组",
        "边数",
        "MTF50",
        "中心MTF50",
        "边角MTF50",
        "MTF50_lpmm",
        "MTF30",
        "MTF10",
        "过锐峰",
        "中心/边",
        "Tenengrad",
        "时间",
        "路径",
        "错误",
    ]
    _xlsx_header(ws, file_headers)
    for r in rows:
        meta = r.get("meta") or {}
        ws.append(
            [
                meta.get("file") or "",
                meta.get("lens") or "",
                _num(meta.get("focal_mm"), 1),
                meta.get("iso") if meta.get("iso") is not None else "",
                meta.get("aperture") or "",
                meta.get("camera") or "",
                r.get("group") or "",
                r.get("n_edges") if r.get("n_edges") is not None else "",
                _num(r.get("mtf50_cy_px")),
                _num(r.get("mtf50_center")),
                _num(r.get("mtf50_corner")),
                _num(r.get("mtf50_lp_mm")),
                _num(r.get("mtf30_cy_px")),
                _num(r.get("mtf10_cy_px")),
                _num(r.get("mtf_peak")),
                _num(r.get("center_edge_ratio"), 3),
                _num(r.get("tenengrad"), 2),
                meta.get("datetime") or "",
                meta.get("path") or r.get("path") or "",
                r.get("error") or "",
            ]
        )
    _xlsx_widths(ws, [22, 28, 8, 8, 8, 16, 28, 8, 10, 11, 11, 12, 10, 10, 10, 10, 12, 18, 40, 16])

    ws2 = wb.create_sheet("分组中位")
    sum_headers = [
        "分组",
        "张数",
        "有效",
        "MTF50中位",
        "MTF50均值",
        "MTF30中位",
        "MTF10中位",
        "过锐峰中位",
        "中心/边中位",
        "MTF50最小",
        "MTF50最大",
    ]
    _xlsx_header(ws2, sum_headers)
    summary = summarize_groups(
        rows,
        "mtf50_cy_px",
        extra=("mtf30_cy_px", "mtf10_cy_px", "mtf_peak", "center_edge_ratio"),
    )
    for s in summary:
        ws2.append(
            [
                s.get("group") or "",
                s.get("n"),
                s.get("n_valid"),
                _num(s.get("median")),
                _num(s.get("mean")),
                _num(s.get("mtf30_cy_px")),
                _num(s.get("mtf10_cy_px")),
                _num(s.get("mtf_peak")),
                _num(s.get("center_edge_ratio"), 3),
                _num(s.get("min")),
                _num(s.get("max")),
            ]
        )
    _xlsx_widths(ws2, [36, 8, 8, 12, 12, 12, 12, 12, 12, 12, 12])
    wb.save(path)


def export_scene_xlsx(path: str, scenes: List[Dict[str, Any]]) -> None:
    from openpyxl import Workbook

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "场景明细"
    file_headers = [
        "场景",
        "名次",
        "相对",
        "镜头",
        "焦距",
        "ISO",
        "光圈",
        "MTF50",
        "MTF30",
        "MTF10",
        "过锐峰",
        "中心/边",
        "Tenengrad",
        "文件",
        "路径",
    ]
    _xlsx_header(ws, file_headers)
    lens_stats: Dict[str, List[float]] = {}
    lens_rank: Dict[str, List[int]] = {}
    for sc in scenes:
        name = sc.get("name")
        for r in sc.get("ranked") or []:
            ws.append(
                [
                    name or "",
                    r.get("rank"),
                    _num(r.get("relative"), 3),
                    r.get("lens") or "",
                    _num(r.get("focal_mm"), 1),
                    r.get("iso") if r.get("iso") is not None else "",
                    r.get("aperture") or "",
                    _num(r.get("mtf50_cy_px")),
                    _num(r.get("mtf30_cy_px")),
                    _num(r.get("mtf10_cy_px")),
                    _num(r.get("mtf_peak")),
                    _num(r.get("center_edge_ratio"), 3),
                    _num(r.get("tenengrad"), 2),
                    r.get("file") or "",
                    r.get("path") or "",
                ]
            )
            lens = str(r.get("lens") or "未知镜头")
            lens_stats.setdefault(lens, []).append(float(r.get("relative") or 0))
            lens_rank.setdefault(lens, []).append(int(r.get("rank") or 0))
    _xlsx_widths(ws, [28, 8, 8, 28, 8, 8, 8, 10, 10, 10, 10, 10, 12, 22, 40])

    ws2 = wb.create_sheet("镜头总评")
    _xlsx_header(ws2, ["镜头", "出场次数", "相对分中位", "平均名次"])
    import numpy as np

    summary = []
    for lens, rels in lens_stats.items():
        arr = np.array(rels, dtype=np.float64)
        ranks = lens_rank[lens]
        summary.append(
            (
                lens,
                len(rels),
                float(np.median(arr)) if arr.size else None,
                float(sum(ranks) / len(ranks)) if ranks else None,
            )
        )
    summary.sort(key=lambda t: -(t[2] or 0))
    for lens, n, med, mean_rank in summary:
        ws2.append([lens, n, _num(med, 3), _num(mean_rank, 2)])
    _xlsx_widths(ws2, [36, 10, 14, 12])
    wb.save(path)


def draw_group_bars(
    ax,
    summary: Sequence[Dict[str, Any]],
    xlabel: str = "MTF50（cycles/pixel）",
    value_key: str = "median",
) -> None:
    prop = cjk_font_prop()
    ax.clear()
    fig = ax.figure
    if not summary:
        ax.set_title("无数据", fontsize=9, pad=6)
        _set_text_font(ax.title, prop)
        return
    items = []
    for s in summary:
        v = s.get(value_key)
        try:
            fv = float(v) if v is not None else 0.0
        except (TypeError, ValueError):
            fv = 0.0
        items.append((str(s.get("group") or ""), fv))
    items.sort(key=lambda x: x[1])
    n = len(items)
    labels = [wrap_group_label(a, max_chars=22) for a, _ in items]
    vals = [b for _, b in items]
    max_lines = max((lab.count("\n") + 1) for lab in labels) if labels else 1
    longest = max(
        (_label_units(line) for lab in labels for line in lab.split("\n")),
        default=8,
    )
    fs = 7.5 if n <= 8 else 6.5
    inch_per = 0.38 + 0.13 * max(0, max_lines - 1)
    fig.set_size_inches(5.8, max(2.8, 0.9 + n * inch_per), forward=True)
    ypos = list(range(n))
    ax.barh(ypos, vals, color="#2E8B57", height=0.55)
    ax.set_yticks(ypos)
    ax.set_yticklabels(labels, fontsize=fs, linespacing=0.95)
    ax.tick_params(axis="y", length=0, pad=3, labelsize=fs)
    ax.tick_params(axis="x", labelsize=7, pad=2)
    ax.set_xlabel(xlabel, fontsize=8, labelpad=4)
    ax.set_title("分组比较（中位数，越高越好）", fontsize=9, pad=6)
    ax.grid(axis="x", linestyle=":", alpha=0.5)
    ax.set_ylim(-0.75, n - 0.25)
    vmax = max(vals) if vals else 1.0
    ax.set_xlim(0, vmax * 1.22 if vmax > 0 else 1.0)
    for i, v in enumerate(vals):
        ax.text(
            v,
            i,
            f" {v:.3f}",
            va="center",
            ha="left",
            fontsize=6.5,
            color="#333333",
            clip_on=False,
        )
    _set_text_font(ax.title, prop)
    _set_text_font(ax.xaxis.label, prop)
    yticks = list(ax.get_yticklabels())
    xticks = list(ax.get_xticklabels())
    for tick in xticks + yticks:
        _set_text_font(tick, prop)
        tick.set_fontsize(fs if tick in yticks else 7)
    left = min(0.60, 0.16 + longest * 0.012)
    fig.subplots_adjust(left=left, right=0.96, top=0.88, bottom=0.18)
    try:
        fig.canvas.draw()
        renderer = fig.canvas.get_renderer()
        fig_w = float(fig.get_window_extent().width)
        max_w = 0.0
        for t in ax.get_yticklabels():
            max_w = max(max_w, float(t.get_window_extent(renderer).width))
        if fig_w > 1:
            left = min(0.64, max(left, (max_w + 18) / fig_w))
            fig.subplots_adjust(left=left, right=0.96, top=0.88, bottom=0.18)
    except Exception:
        pass
