# -*- coding: utf-8 -*-
"""ISO 12233 斜边 MTF 与分组比较（无 Qt）。"""

from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
TOOL = ROOT / "tools" / "MTF-meter"
if str(TOOL) not in sys.path:
    sys.path.insert(0, str(TOOL))

from mtf_meter.exif_meta import group_label, read_photo_meta  # noqa: E402
from mtf_meter.grouping import (  # noqa: E402
    attach_group,
    cluster_label,
    cluster_scenes_by_folder,
    cluster_scenes_by_lens,
    cluster_scenes_by_time,
    summarize_groups,
)
from mtf_meter.image_load import collect_images  # noqa: E402
from mtf_meter.relative import rank_scene  # noqa: E402
from mtf_meter.slanted_edge import (  # noqa: E402
    make_slanted_edge,
    slanted_edge_sfr,
    theoretical_gaussian_mtf50,
)


def test_gaussian_mtf50_near_theory():
    sigma = 1.3
    img = make_slanted_edge(h=200, w=200, angle_deg=5.0, sigma=sigma)
    got = slanted_edge_sfr(img, oversample=4)
    assert got is not None
    expect = theoretical_gaussian_mtf50(sigma)
    # 斜边 SFR 带窗与离散化，绝对值允许偏差；须在合理量级且低于奈奎斯特
    assert 0.5 * expect < got["mtf50_cy_px"] < 2.0 * expect
    assert got["mtf50_cy_px"] < 0.5
    assert got["contrast"] > 50


def test_sharper_edge_higher_mtf50():
    sharp = slanted_edge_sfr(make_slanted_edge(sigma=0.7))
    soft = slanted_edge_sfr(make_slanted_edge(sigma=2.2))
    assert sharp is not None and soft is not None
    assert sharp["mtf50_cy_px"] > soft["mtf50_cy_px"]


def test_center_edge_ratio_and_extra_metrics():
    from mtf_meter.analyze import center_edge_ratio

    assert abs(center_edge_ratio(0.24, 0.12) - 2.0) < 1e-9
    assert center_edge_ratio(0.24, None) is None
    assert center_edge_ratio(0.24, 0.0) is None
    rows = [
        {
            "meta": {"lens": "A", "focal_mm": 400.0, "iso": 800},
            "mtf50_cy_px": 0.22,
            "mtf30_cy_px": 0.31,
            "mtf_peak": 1.12,
            "center_edge_ratio": 1.4,
        },
        {
            "meta": {"lens": "A", "focal_mm": 400.0, "iso": 800},
            "mtf50_cy_px": 0.24,
            "mtf30_cy_px": 0.33,
            "mtf_peak": 1.04,
            "center_edge_ratio": 1.2,
        },
    ]
    attach_group(rows, ("lens", "focal"))
    summary = summarize_groups(
        rows, "mtf50_cy_px", extra=("mtf30_cy_px", "mtf_peak", "center_edge_ratio")
    )
    rec = summary[0]
    assert abs(rec["mtf30_cy_px"] - 0.32) < 1e-9
    assert abs(rec["mtf_peak"] - 1.08) < 1e-9
    assert abs(rec["center_edge_ratio"] - 1.3) < 1e-9


def test_export_iso_xlsx_includes_group_medians(tmp_path):
    from mtf_meter.export import export_iso_xlsx
    from openpyxl import load_workbook

    rows = [
        {
            "meta": {
                "file": "a.jpg",
                "lens": "A",
                "focal_mm": 400,
                "iso": 200,
                "path": "a.jpg",
            },
            "group": "A · 400mm",
            "n_edges": 4,
            "mtf50_cy_px": 0.20,
            "mtf50_center": 0.22,
            "mtf50_corner": 0.16,
            "mtf30_cy_px": 0.28,
            "mtf10_cy_px": 0.36,
            "mtf_peak": 1.11,
            "center_edge_ratio": 1.375,
        },
        {
            "meta": {
                "file": "b.jpg",
                "lens": "A",
                "focal_mm": 400,
                "iso": 200,
                "path": "b.jpg",
            },
            "group": "A · 400mm",
            "n_edges": 3,
            "mtf50_cy_px": 0.24,
            "mtf50_center": 0.26,
            "mtf50_corner": 0.18,
            "mtf30_cy_px": 0.32,
            "mtf10_cy_px": 0.40,
            "mtf_peak": 1.05,
            "center_edge_ratio": 1.25,
        },
    ]
    out = tmp_path / "iso.xlsx"
    export_iso_xlsx(str(out), rows)
    wb = load_workbook(out)
    assert wb.sheetnames == ["文件明细", "分组中位"]
    files = wb["文件明细"]
    assert files["A1"].value == "文件"
    assert files.max_row == 3
    med = wb["分组中位"]
    assert med["A1"].value == "分组"
    assert med["D1"].value == "MTF50中位"
    assert med.max_row == 2
    assert med["A2"].value == "A · 400mm"
    assert abs(float(med["D2"].value) - 0.22) < 1e-9


def test_export_scene_xlsx_has_lens_sheet(tmp_path):
    from mtf_meter.export import export_scene_xlsx
    from openpyxl import load_workbook

    pack = [
        {
            "name": "场景01",
            "ranked": [
                {
                    "rank": 1,
                    "relative": 1.0,
                    "lens": "sharp",
                    "focal_mm": 400,
                    "mtf50_cy_px": 0.2,
                    "file": "h.jpg",
                    "path": "h.jpg",
                },
                {
                    "rank": 2,
                    "relative": 0.5,
                    "lens": "soft",
                    "focal_mm": 400,
                    "mtf50_cy_px": 0.1,
                    "file": "s.jpg",
                    "path": "s.jpg",
                },
            ],
        }
    ]
    out = tmp_path / "scene.xlsx"
    export_scene_xlsx(str(out), pack)
    wb = load_workbook(out)
    assert "场景明细" in wb.sheetnames
    assert "镜头总评" in wb.sheetnames
    lens = wb["镜头总评"]
    names = {lens.cell(i, 1).value for i in range(2, lens.max_row + 1)}
    assert names == {"sharp", "soft"}


def test_suggested_xlsx_path_uses_folder_date_and_number(tmp_path):
    from datetime import datetime

    from mtf_meter.export import suggested_xlsx_path

    folder = tmp_path / "RF100-500_ISO卡"
    folder.mkdir()
    now = datetime(2026, 9, 11, 12, 0, 0)
    p1 = suggested_xlsx_path(str(folder), now=now)
    assert p1.parent == folder
    assert p1.name == "RF100-500_ISO卡_20260911_01.xlsx"
    p1.write_bytes(b"x")
    p2 = suggested_xlsx_path(str(folder), now=now)
    assert p2.name == "RF100-500_ISO卡_20260911_02.xlsx"


def test_short_group_label_single_line():
    from mtf_meter.export import short_group_label

    s = short_group_label("FE 200-600mm F5.6-6.3 G OSS · 600mm · ISO800", 16)
    assert "\n" not in s
    assert s.endswith("…")
    assert len(s) <= 16


def test_export_iso_csv_includes_mtf30_peak_ratio(tmp_path):
    from mtf_meter.export import export_iso_csv

    out = tmp_path / "iso.csv"
    export_iso_csv(
        str(out),
        [
            {
                "meta": {"file": "a.jpg", "lens": "A", "focal_mm": 400, "iso": 200},
                "n_edges": 4,
                "mtf50_cy_px": 0.20,
                "mtf50_center": 0.22,
                "mtf50_corner": 0.16,
                "mtf30_cy_px": 0.28,
                "mtf10_cy_px": 0.36,
                "mtf_peak": 1.11,
                "center_edge_ratio": 1.375,
            }
        ],
    )
    text = out.read_text(encoding="utf-8-sig")
    assert "mtf30_cy_px" in text
    assert "mtf10_cy_px" in text
    assert "mtf_peak" in text
    assert "center_edge_ratio" in text
    assert "1.1100" in text or "1.11" in text


def test_rank_scene_keeps_sfr_extras():
    rows = [
        {
            "meta": {"lens": "soft", "file": "s.jpg"},
            "mtf50_cy_px": 0.10,
            "mtf30_cy_px": 0.14,
            "mtf10_cy_px": 0.20,
            "mtf_peak": 1.02,
            "center_edge_ratio": 1.3,
        },
        {
            "meta": {"lens": "sharp", "file": "h.jpg"},
            "mtf50_cy_px": 0.20,
            "mtf30_cy_px": 0.28,
            "mtf10_cy_px": 0.36,
            "mtf_peak": 1.15,
            "center_edge_ratio": 1.1,
        },
    ]
    ranked = rank_scene(rows)
    by = {r["lens"]: r for r in ranked}
    assert by["sharp"]["mtf30_cy_px"] == 0.28
    assert by["sharp"]["mtf10_cy_px"] == 0.36
    assert by["sharp"]["mtf_peak"] == 1.15
    assert by["soft"]["center_edge_ratio"] == 1.3


def test_group_label_and_summary():
    rows = [
        {
            "meta": {"lens": "A", "focal_mm": 400.0, "iso": 800},
            "mtf50_cy_px": 0.22,
        },
        {
            "meta": {"lens": "A", "focal_mm": 400.0, "iso": 800},
            "mtf50_cy_px": 0.24,
        },
        {
            "meta": {"lens": "B", "focal_mm": 400.0, "iso": 800},
            "mtf50_cy_px": 0.18,
        },
    ]
    attach_group(rows, ("lens", "focal"))
    assert rows[0]["group"] == group_label(rows[0]["meta"], ("lens", "focal"))
    summary = summarize_groups(rows, "mtf50_cy_px")
    by = {s["group"]: s for s in summary}
    assert by[rows[0]["group"]]["n"] == 2
    assert abs(by[rows[0]["group"]]["median"] - 0.23) < 1e-9
    assert by[rows[2]["group"]]["median"] < by[rows[0]["group"]]["median"]


def test_cluster_scenes_by_time():
    t0 = datetime(2026, 9, 3, 10, 0, 0).timestamp()
    rows = [
        {"meta": {"datetime_ts": t0, "path": "a1.jpg"}},
        {"meta": {"datetime_ts": t0 + 30, "path": "a2.jpg"}},
        {"meta": {"datetime_ts": t0 + 900, "path": "b1.jpg"}},
    ]
    cl = cluster_scenes_by_time(rows, gap_sec=180)
    assert len(cl) == 2
    assert len(cl[0]) == 2
    assert len(cl[1]) == 1


def test_cluster_scenes_by_folder(tmp_path):
    rows = [
        {"meta": {"path": str(tmp_path / "lensA" / "1.jpg")}},
        {"meta": {"path": str(tmp_path / "lensA" / "2.jpg")}},
        {"meta": {"path": str(tmp_path / "lensB" / "1.jpg")}},
    ]
    cl = cluster_scenes_by_folder(rows)
    sizes = sorted(len(c) for c in cl)
    assert sizes == [1, 2]


def test_cluster_scenes_by_lens():
    rows = [
        {"meta": {"lens": "RF 100-500", "path": "a1.jpg"}},
        {"meta": {"lens": "RF 100-500", "path": "a2.jpg"}},
        {"meta": {"lens": "RF 600", "path": "b1.jpg"}},
        {"meta": {"lens": "", "path": "u.jpg"}},
    ]
    cl = cluster_scenes_by_lens(rows)
    sizes = sorted(len(c) for c in cl)
    assert sizes == [1, 1, 2]
    labels = {cluster_label(c, "lens", i) for i, c in enumerate(cl, start=1)}
    assert any("RF 100-500" in x and "2张" in x for x in labels)
    assert any("未知镜头" in x for x in labels)


def test_rank_scene_relative():
    rows = [
        {"meta": {"lens": "soft", "file": "s.jpg"}, "mtf50_cy_px": 0.10},
        {"meta": {"lens": "sharp", "file": "h.jpg"}, "mtf50_cy_px": 0.20},
    ]
    ranked = rank_scene(rows)
    assert ranked[0]["lens"] == "sharp"
    assert abs(ranked[0]["relative"] - 1.0) < 1e-9
    assert abs(ranked[1]["relative"] - 0.5) < 1e-9
    assert "edges" in ranked[0]


def test_overlay_rois_draws_center_box():
    from mtf_meter.preview import overlay_rois

    img = np.full((80, 120, 3), 30, dtype=np.uint8)
    vis = overlay_rois(
        img,
        [
            {
                "roi": [10, 15, 50, 55],
                "zone": "center",
                "mtf50_cy_px": 0.22,
                "angle_deg": 5.0,
            }
        ],
    )
    assert vis.shape[0] == 80 and vis.shape[1] == 120
    assert int(vis[:, :, 1].max()) > 100


def test_measure_image_keeps_roi_and_angle(tmp_path):
    from mtf_meter.analyze import measure_image

    gray = make_slanted_edge(h=240, w=240, angle_deg=5.0, sigma=1.0)
    bgr = np.stack([gray, gray, gray], axis=-1).astype(np.uint8)
    p = tmp_path / "chart.jpg"
    cv2.imwrite(str(p), bgr)
    rec = measure_image(str(p))
    assert rec["ok"]
    assert rec["n_edges"] >= 1
    e0 = rec["edges"][0]
    assert len(e0["roi"]) == 4
    assert e0["roi"][2] > e0["roi"][0]
    assert "angle_deg" in e0


def test_collect_images_skips_dot_dirs(tmp_path):
    (tmp_path / "ok.jpg").write_bytes(b"x")
    hidden = tmp_path / ".thumb"
    hidden.mkdir()
    (hidden / "x.jpg").write_bytes(b"x")
    # 真实可读图
    img = np.zeros((8, 8, 3), dtype=np.uint8)
    cv2.imwrite(str(tmp_path / "a.jpg"), img)
    found = collect_images(str(tmp_path))
    names = [Path(p).name for p in found]
    assert "a.jpg" in names
    assert "x.jpg" not in names


def test_wrap_group_label_breaks_on_dot_and_length():
    from mtf_meter.export import wrap_group_label

    s = wrap_group_label("FE 200-600mm F5.6-6.3 G OSS · 600mm · ISO800")
    assert "\n" in s
    assert "600mm" in s.split("\n")
    long = "甲" * 40
    w = wrap_group_label(long, max_chars=18)
    assert w.count("\n") >= 1
    assert all(len(line) <= 18 for line in w.split("\n"))


def test_draw_group_bars_chinese_labels():
    import matplotlib

    try:
        matplotlib.use("Agg")
    except Exception:
        pass
    from matplotlib.figure import Figure

    from mtf_meter.export import cjk_font_prop, draw_group_bars

    prop = cjk_font_prop()
    fig = Figure(figsize=(4, 2))
    ax = fig.add_subplot(111)
    draw_group_bars(
        ax,
        [
            {"group": "未知镜头 · 400mm", "median": 0.21},
            {"group": "分组比较测试", "median": 0.18},
        ],
        xlabel="MTF50（cycles/pixel）",
    )
    labels = [t.get_text() for t in ax.get_yticklabels()]
    joined = " ".join(labels)
    assert "未知镜头" in joined
    assert "400mm" in joined
    assert any("\n" in x for x in labels)
    title = ax.get_title()
    assert "分组比较" in title
    if prop is not None:
        assert ax.title.get_fontname()


def test_read_photo_meta_without_exif(tmp_path):
    img = np.zeros((16, 16, 3), dtype=np.uint8)
    p = tmp_path / "n.jpg"
    cv2.imwrite(str(p), img)
    meta = read_photo_meta(str(p))
    assert meta["file"] == "n.jpg"
    assert "lens" in meta


def test_find_edge_rois_four_slots_stable():
    from mtf_meter.edges import find_edge_rois

    canvas = np.full((520, 520), 90.0)
    patch = make_slanted_edge(h=110, w=110, angle_deg=5.0, sigma=1.0)
    # 与四区窗口对齐：左、右、上、中下（不要放几何中心/四角）
    places = [(200, 8), (200, 400), (20, 205), (290, 205)]
    for y, x in places:
        canvas[y : y + 110, x : x + 110] = patch
    g0 = np.clip(canvas, 0, 255).astype(np.uint8)
    rois = find_edge_rois(g0)
    assert 1 <= len(rois) <= 4
    slots = [r["slot"] for r in rois]
    assert len(slots) == len(set(slots))
    assert set(slots).issubset({"left", "right", "top", "center"})
    sides = [r["x2"] - r["x1"] for r in rois]
    assert max(sides) - min(sides) <= 2
    rng = np.random.RandomState(3)
    g1 = np.clip(g0.astype(np.float32) + rng.normal(0, 1.2, g0.shape), 0, 255).astype(
        np.uint8
    )
    rois2 = find_edge_rois(g1)
    assert {r["slot"] for r in rois} == {r["slot"] for r in rois2}


def test_find_edge_rois_never_more_than_four():
    from mtf_meter.edges import find_edge_rois

    rng = np.random.RandomState(9)
    noise = rng.randint(0, 255, (360, 480), dtype=np.uint8)
    assert len(find_edge_rois(noise)) <= 4


def test_find_edge_rois_skips_center_rings_and_corner_stars():
    from mtf_meter.edges import find_edge_rois

    h = w = 640
    canvas = np.full((h, w), 110.0)
    yy, xx = np.ogrid[:h, :w]
    # 正中圆环放在上区与中下区之间的空隙，旧逻辑会当成 C
    cr, cc = 275, 320
    rr = 40
    rings = ((np.hypot(xx - cc, yy - cr).astype(int) // 6) % 2) * 200 + 30
    canvas[cr - rr : cr + rr, cc - rr : cc + rr] = rings[
        cr - rr : cr + rr, cc - rr : cc + rr
    ]
    # 左上星形图：旧逻辑会当成 TL
    sy, sx = 70, 70
    star = ((np.arctan2(yy - sy, xx - sx) * 10 / np.pi).astype(int) % 2) * 200 + 30
    canvas[10:130, 10:130] = star[10:130, 10:130]
    patch = make_slanted_edge(h=120, w=120, angle_deg=5.0, sigma=1.0)
    # 中下斜方块（应成为 C）
    canvas[360:480, 260:380] = patch
    # 左右黑斜块
    canvas[240:360, 10:130] = patch
    canvas[240:360, 510:630] = patch
    g = np.clip(canvas, 0, 255).astype(np.uint8)
    rois = find_edge_rois(g)
    slots = {r["slot"] for r in rois}
    assert "tl" not in slots and "tr" not in slots
    assert "br" not in slots and "bottom" not in slots
    assert "left" in slots
    assert "right" in slots
    assert "center" in slots
    center = next(r for r in rois if r["slot"] == "center")
    cy = 0.5 * (center["y1"] + center["y2"])
    # 框应落在中下斜块，而不是正中圆环
    assert cy > 300


def test_find_edge_rois_highres_still_finds():
    from mtf_meter.edges import find_edge_rois

    # 直接生成大图（不要把 200px 硬拉到 2400，Canny 会糊到找不到）
    big = np.clip(
        make_slanted_edge(h=1600, w=1600, angle_deg=5.0, sigma=2.0), 0, 255
    ).astype(np.uint8)
    rois = find_edge_rois(big)
    assert len(rois) >= 1
    assert len(rois) <= 4


def test_measure_image_on_synthetic_chart(tmp_path):
    from mtf_meter.analyze import measure_image

    gray = make_slanted_edge(h=240, w=240, angle_deg=5.0, sigma=1.0)
    bgr = np.stack([gray, gray, gray], axis=-1).astype(np.uint8)
    p = tmp_path / "chart.jpg"
    cv2.imwrite(str(p), bgr)
    rec = measure_image(str(p))
    assert rec["ok"]
    assert rec["n_edges"] >= 1
    assert rec["n_edges"] <= 4
    assert rec["mtf50_cy_px"] is not None
    assert rec["mtf30_cy_px"] is not None
    assert rec["mtf10_cy_px"] is not None
    assert rec["mtf_peak"] is not None
    assert 0.05 < float(rec["mtf50_cy_px"]) < 0.48
    assert float(rec["mtf30_cy_px"]) >= float(rec["mtf50_cy_px"]) - 1e-6
    assert "center_edge_ratio" in rec


def test_measure_many_emits_row_and_can_stop(tmp_path):
    from mtf_meter.analyze import measure_many

    gray = make_slanted_edge(h=160, w=160, angle_deg=5.0, sigma=1.0)
    bgr = np.stack([gray, gray, gray], axis=-1).astype(np.uint8)
    paths = []
    for i in range(3):
        p = tmp_path / f"c{i}.jpg"
        cv2.imwrite(str(p), bgr)
        paths.append(str(p))
    seen = []
    state = {"n": 0}

    def prog(d):
        if d.get("kind") == "row":
            seen.append(d)
            state["n"] = int(d["done"])
            assert d.get("row", {}).get("ok") is True

    def should():
        return state["n"] >= 1

    rows = measure_many(paths, progress=prog, should_cancel=should)
    assert len(rows) == 1
    assert len(seen) == 1
    assert seen[0]["total"] == 3
    assert seen[0]["done"] == 1
